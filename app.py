import streamlit as st
import pandas as pd
import openpyxl, io, zipfile, os, tempfile
from pathlib import Path

st.set_page_config(page_title='Connect Sherpur GA Auto Report', page_icon='📊', layout='centered')
st.title('📊 Connect Sherpur GA — Auto Report')
st.write('ফোন থেকেই ডাউনলোড করা রিপোর্ট ফাইল আপলোড করে Excel রিপোর্ট তৈরি করুন।')
st.info('সমর্থিত: XLS/XLSX, CSV, ট্যাব-ডিলিমিটেড রিপোর্ট, ZIP (ভেতরে সমর্থিত ফাইল), এবং টেক্সট-ভিত্তিক PDF।')

TEMPLATE = Path(__file__).with_name('template.xlsx')
EXPECTED = ['CLUSTER_NAME','REGION_NAME','ACTIVATION_DATE','DISTRIBUTOR_CODE','DISTRIBUTOR_NAME','RETAILER_CODE','RETAILER_NAME','PRODUCT_NAME','PRODUCT_CODE','PROMOTION','SELLING_PRICE','Qty','BP_FLAG','BP_NUMBER']

def read_one(name, data):
    ext = Path(name).suffix.lower()
    if ext in ['.xlsx', '.xlsm']:
        df = pd.read_excel(io.BytesIO(data), dtype=str)
    elif ext == '.xls':
        # Some activation exports use tab-delimited text despite the .xls extension.
        try:
            df = pd.read_csv(io.BytesIO(data), sep='\t', dtype=str, encoding='utf-8-sig')
            if len(df.columns) < 3:
                raise ValueError('not tab-separated')
        except Exception:
            df = pd.read_excel(io.BytesIO(data), dtype=str, engine='xlrd')
    elif ext == '.csv':
        try: df = pd.read_csv(io.BytesIO(data), dtype=str, encoding='utf-8-sig')
        except Exception: df = pd.read_csv(io.BytesIO(data), dtype=str, encoding='latin1')
    elif ext == '.txt':
        df = pd.read_csv(io.BytesIO(data), sep=None, engine='python', dtype=str, encoding='utf-8-sig')
    elif ext == '.pdf':
        try:
            import pdfplumber
            table_frames=[]
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                for page in pdf.pages:
                    for table in (page.extract_tables() or []):
                        if table and len(table) >= 2:
                            header=[str(x or '').strip() for x in table[0]]
                            body=table[1:]
                            if sum(bool(x) for x in header) >= 4:
                                table_frames.append(pd.DataFrame(body, columns=header))
            if not table_frames:
                raise ValueError('PDF-এ ব্যবহারযোগ্য টেবিল পাওয়া যায়নি। স্ক্যান করা PDF হলে OCR প্রয়োজন; সম্ভব হলে মূল Excel/CSV আপলোড করুন।')
            df=pd.concat(table_frames, ignore_index=True).dropna(how='all')
        except ImportError:
            raise ValueError('PDF পড়ার জন্য প্রয়োজনীয় লাইব্রেরি ইনস্টল নেই।')
    else:
        raise ValueError(f'অসমর্থিত ফাইল: {name}')
    df.columns=[str(c).strip() for c in df.columns]
    return df

def collect_upload(upload):
    files=[]
    name=upload.name
    data=upload.getvalue()
    if Path(name).suffix.lower()=='.zip':
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for n in z.namelist():
                if n.endswith('/') or n.startswith('__MACOSX/'): continue
                if Path(n).suffix.lower() in ['.xls','.xlsx','.xlsm','.csv','.txt','.pdf']:
                    files.append((n,z.read(n)))
    else: files=[(name,data)]
    frames=[]
    for n,b in files:
        df=read_one(n,b)
        if 'PDF_TEXT' in df.columns: continue
        frames.append((n,df))
    if not frames: raise ValueError('ZIP-এর মধ্যে ব্যবহারযোগ্য Excel/CSV/Text ফাইল পাওয়া যায়নি।')
    # Merge only files with activation-report style headers; avoid accidental mixing of unrelated tables.
    good=[]
    for n,df in frames:
        norm={c.upper().replace(' ','_'):c for c in df.columns}
        required=['CLUSTER_NAME','REGION_NAME','ACTIVATION_DATE','DISTRIBUTOR_CODE','DISTRIBUTOR_NAME','RETAILER_CODE','RETAILER_NAME','PRODUCT_NAME','PRODUCT_CODE','PROMOTION','SELLING_PRICE','QTY','BP_FLAG']
        if all(k in norm for k in required):
            rename={norm[k]:k for k in required}
            if 'BP_NUMBER' in norm: rename[norm['BP_NUMBER']]='BP_NUMBER'
            df=df.rename(columns=rename)
            if 'BP_NUMBER' not in df: df['BP_NUMBER']=''
            good.append(df[EXPECTED])
    if not good:
        first=frames[0][1]
        raise ValueError('আপলোড করা ফাইলে ActivationReport-এর প্রয়োজনীয় কলামগুলো পাওয়া যায়নি। প্রথম ফাইলের কলাম: '+', '.join(map(str,first.columns[:25])))
    return pd.concat(good,ignore_index=True).fillna('')

def clear_and_write(ws, start_col, rows, max_cols=14, start_row=2):
    # Clear the helper-data block only; report formulas and target tables are kept.
    end_col=start_col+max_cols-1
    for r in range(start_row, max(ws.max_row if ws.max_row < 10000 else 5000, start_row)+1):
        for c in range(start_col,end_col+1):
            cell=ws.cell(r,c)
            if cell.data_type != 'f': cell.value=None
    for idx, row in enumerate(rows, start=start_row):
        for j, val in enumerate(row, start=start_col):
            if isinstance(val, str) and val.strip().lower() in ('nan','nat','none'): val=''
            # Convert quantity/price to numbers when possible
            if j in (start_col+10,start_col+11):
                try: val=float(val) if val!='' else None
                except Exception: pass
            ws.cell(idx,j).value = val

def row_values(df):
    vals=[]
    for _,r in df.iterrows():
        vals.append([r.get(c,'') for c in EXPECTED])
    return vals

uploaded=st.file_uploader('ডাউনলোড করা রিপোর্ট ফাইল নির্বাচন করুন', type=['xls','xlsx','xlsm','csv','txt','pdf','zip'])
if uploaded:
    st.write(f'নির্বাচিত ফাইল: **{uploaded.name}**')
    if st.button('⚙️ GENERATE REPORT', type='primary', use_container_width=True):
        try:
            df=collect_upload(uploaded)
            st.success(f'ফাইল পড়া হয়েছে: {len(df):,}টি রেকর্ড')
            if not TEMPLATE.exists(): raise ValueError('মূল রিপোর্ট টেমপ্লেট পাওয়া যায়নি।')
            # Load formula template; set Excel to recalculate formulas when opened.
            wb=openpyxl.load_workbook(TEMPLATE)
            df['PRODUCT_CODE']=df['PRODUCT_CODE'].astype(str).str.strip().str.upper()
            daily_mmst=df[df['PRODUCT_CODE'].eq('MMST')]
            daily_mmstc=df[df['PRODUCT_CODE'].eq('MMSTC')]
            all_rows=row_values(df)
            mmst_rows=row_values(daily_mmst)
            mmstc_rows=row_values(daily_mmstc)
            # Daily and monthly raw-data helper blocks.
            for sn in ['TC GA','Monthly GAt']:
                ws=wb[sn]
                clear_and_write(ws,14,mmst_rows,14)  # N:AA, MMST
                clear_and_write(ws,30,mmstc_rows,14) # AD:AQ, MMSTC
            for sn in [' GA ','Monthly GA']:
                ws=wb[sn]
                clear_and_write(ws,12,all_rows,14) # L:Y
            try:
                wb.calculation.fullCalcOnLoad=True
                wb.calculation.forceFullCalc=True
                wb.calculation.calcMode='auto'
            except Exception: pass
            out=io.BytesIO(); wb.save(out); out.seek(0)
            st.success('রিপোর্ট ফাইল তৈরি হয়েছে। Excel-এ খুললে ফর্মুলাগুলো পুনরায় হিসাব হবে।')
            st.download_button('📥 DOWNLOAD REPORT (XLSX)', data=out.getvalue(), file_name='Connect_Sherpur_GA_Report.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet', use_container_width=True)
            st.caption(f'ইনপুট রেকর্ড: {len(df):,} | MMST: {len(daily_mmst):,} | MMSTC: {len(daily_mmstc):,}')
        except Exception as e:
            st.error(str(e))

with st.expander('ব্যবহারের নিয়ম'):
    st.markdown('''1. ডাউনলোড করা Activation Report আপলোড করুন। ZIP হলে তার ভেতরে Excel/CSV/TXT ফাইল রাখুন।\n2. **GENERATE REPORT** চাপুন।\n3. তৈরি হওয়া XLSX ফাইল ডাউনলোড করুন।\n4. প্রথমবার ফলাফল যাচাই করুন। আপনার ফাইলে যদি অন্য ধরনের কলাম/রিপোর্ট থাকে, সেগুলোর জন্য আলাদা mapping লাগতে পারে।''')
